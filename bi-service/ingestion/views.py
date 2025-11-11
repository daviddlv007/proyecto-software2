# =============================================
# 📦 Django
# =============================================
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST, require_GET
from django.views.decorators.csrf import csrf_exempt
from django.core.files.storage import default_storage
from django.core.serializers.json import DjangoJSONEncoder
from django.http import JsonResponse, HttpResponse, FileResponse
from django.urls import reverse
from django.conf import settings
from django.db import connection
from django.core.mail import EmailMessage

# =============================================
#  Interno (modelos y servicios)
# =============================================
from .models import DataSource, UploadedDataset, Diagrama, ExternalConnection
from .services import (
    import_csv_or_excel,
    import_sql_script,
    sanitize_identifier,
    get_dataset,
    get_schema_info,
    generar_consulta_y_grafico,
    generar_diagramas_automaticos,
    generar_diagrama_chat,
    guardar_diagrama,
    obtener_diagramas_por_archivo,
    generar_chat_chart_o_prediccion,
    get_engine,
    ejecutar_sql_para_chart,
    reduce_schema,
    analyze_chart_image,
)

# =============================================
# ⚙️ Librerías externas
# =============================================
from sqlalchemy import inspect, text as sa_text 
import subprocess
import tempfile
import json
import uuid
import os
import io
import pandas as pd
import re
import base64
from datetime import datetime
from openpyxl.utils import get_column_letter
from json import JSONDecodeError

@login_required
def upload_dataset_view(request):
    """
    Soporta:
    ✅ Subir archivo (CSV/XLSX/SQL).
    ✅ Conectar a PostgreSQL externa (host/port/db/user/password).
    ✅ Genera diagramas automáticos para cada tabla detectada.
    """
    if request.method == "POST":
        source_type = request.POST.get("source_type", "file")

        # ======================================================
        # 📁 1) IMPORTAR ARCHIVO (CSV / XLSX / SQL)
        # ======================================================
        if source_type == "file" and request.FILES.get("file"):
            f = request.FILES["file"]
            ext = f.name.split(".")[-1].lower()

            if ext not in ("csv", "xlsx", "sql"):
                return render(request, "ingestion/upload.html", {"error": "Solo CSV, XLSX o SQL"})

            schema = sanitize_identifier(f"user_{request.user.id}_file_{uuid.uuid4().hex[:8]}")
            table = sanitize_identifier(f"ds_{uuid.uuid4().hex[:6]}")

            ds = DataSource.objects.create(
                name=request.POST.get("name", f.name),
                kind=DataSource.FILE,
                owner=request.user,
                internal_schema=schema,
                internal_table=table,
            )

            up = UploadedDataset.objects.create(source=ds, file=f, file_type=ext)
            path = up.file.path

            try:
                if ext in ("csv", "xlsx"):
                    meta = import_csv_or_excel(path, schema, table)
                    up.rows_ingested = meta.get("rows", 0)
                    up.columns = meta.get("columns", [])
                    up.save(update_fields=["rows_ingested", "columns"])
                else:
                    tables, meta_info, main_table = import_sql_script(path, schema)
                    ds.internal_table = main_table or (tables[0] if tables else table)
                    ds.save(update_fields=["internal_table"])

            except Exception as e:
                ds.delete()
                return render(request, "ingestion/upload.html", {"error": f"Error importando archivo: {e}"})

            engine = get_engine(**ds.get_db_credentials())
            insp = inspect(engine)
            existing_tables = insp.get_table_names(schema=ds.internal_schema or "public")

            if not ds.internal_table or ds.internal_table not in existing_tables:
                return render(request, "ingestion/upload.html", {
                    "error": f"No se encontró la tabla '{ds.internal_table}' en el esquema '{ds.internal_schema}'."
                })

            try:
                generar_diagramas_automaticos(ds, **ds.get_db_credentials())
            except Exception as e:
                print("Error generando diagramas automáticos:", e)

            return redirect("ingestion:list")

        # ======================================================
        # 🟢 2) CONEXIÓN A POSTGRESQL EXTERNA
        # ======================================================
        if source_type == "db":
            host = request.POST.get("db_host", "").strip() or "postgres"
            port = int(request.POST.get("db_port") or 5432)  # 👈 por defecto 5432 para Docker
            database = request.POST.get("db_name", "").strip()
            user = request.POST.get("db_user", "").strip()
            password = request.POST.get("db_password", "").strip()
            schema = request.POST.get("db_schema", "").strip() or "public"
            table = request.POST.get("db_table", "").strip() or None

            if not (host and database and user and password):
                return render(request, "ingestion/upload.html", {
                    "error": "Rellena host, BD, usuario y contraseña para conectar."
                })

            # ✅ Crear DataSource con credenciales
            ds = DataSource.objects.create(
                name=request.POST.get("name", f"external_{user}@{host}:{port}/{database}"),
                kind=DataSource.LIVE,
                owner=request.user,
                internal_schema=sanitize_identifier(schema),
                internal_table=sanitize_identifier(table) if table else "",
                db_host=host,
                db_port=port,
                db_name=database,
                db_user=user,
                db_password=password,
            )

            # ✅ Guardar conexión cifrada (para uso posterior por get_db_credentials)
            ExternalConnection.objects.create(
                source=ds,
                db_type=ExternalConnection.POSTGRES,
                host=host,
                port=port,
                database=database,
                username=user,
                password=password,
                extras={"schema": schema},
            )

            creds = ds.get_db_credentials()

            # Test conexión
            try:
                engine = get_engine(**creds)
                with engine.begin() as conn:
                    conn.execute(sa_text("SELECT 1"))
            except Exception as e:
                ds.delete()
                return render(request, "ingestion/upload.html", {"error": f"No se pudo conectar: {e}"})

            # === Obtener tablas y generar diagramas automáticos ===
            engine = get_engine(**creds)
            insp = inspect(engine)
            tablas = insp.get_table_names(schema=schema)

            if not tablas:
                ds.delete()
                return render(request, "ingestion/upload.html", {
                    "error": f"No se encontraron tablas dentro del esquema '{schema}'."
                })

            print(f"🔍 Analizando BD externa — esquema '{schema}', {len(tablas)} tablas")

            for t in tablas:
                ds.internal_table = t
                ds.save(update_fields=["internal_table"])
                try:
                    generar_diagramas_automaticos(ds, **creds)
                    print(f"✅ Diagramas generados para '{t}'")
                except Exception as e:
                    print(f"⚠️ Error generando diagramas para '{t}': {e}")

            return redirect("ingestion:list")

    return render(request, "ingestion/upload.html")


@login_required
def list_sources_view(request):
    sources = DataSource.objects.filter(owner=request.user).order_by('-created_at')

    for source in sources:
        source.tables = []
        if not source.internal_schema:
            continue
        try:
            creds = source.get_db_credentials() or {}
            # si quieres fijar schema para ese engine (no obligatorio aquí):
            engine = get_engine(**creds)
            with engine.connect() as conn:
                result = conn.execute(sa_text("""
                    SELECT table_name
                    FROM information_schema.tables
                    WHERE table_schema = :schema
                    ORDER BY table_name
                """), {"schema": source.internal_schema})
                source.tables = [row[0] for row in result]
        except Exception as e:
            print(f"Error obteniendo tablas para {source.name}: {e}")
            source.tables = []

    return render(request, 'ingestion/list.html', {'sources': sources})
from django.conf import settings

def normalize_db_credentials(db_credentials: dict, data_source=None) -> dict:
    """
    Normaliza credenciales para psycopg2/sqlalchemy.
    Usa `database` en lugar de `dbname`.
    """
    if db_credentials:
        return {
            "host": db_credentials.get("host", "postgres"),
            "database": db_credentials.get("database") or (data_source.db_name if data_source else None),
            "user": db_credentials.get("user") or (data_source.db_user if data_source else None),
            "password": db_credentials.get("password") or (data_source.db_password if data_source else None),
            "port": db_credentials.get("port") or (data_source.db_port if data_source else 5432),
            "schema": db_credentials.get("schema", "public"),
        }

    db = settings.DATABASES["default"]
    return {
        "host": db.get("HOST", "postgres"),
        "database": db.get("NAME"),
        "user": db.get("USER"),
        "password": db.get("PASSWORD"),
        "port": db.get("PORT", 5432),
        "schema": "public",
    }



def chart_view(request, source_id):
    source = DataSource.objects.get(id=source_id)
    df = get_dataset(source.internal_schema, source.internal_table, **source.get_db_credentials())

    if "fecha" in df.columns:
        df["fecha"] = df["fecha"].astype(str)

    # Convertir a lista de diccionarios
    chart_data = df.to_dict(orient="records")

    # Convertir a JSON válido
    chart_data_json = json.dumps(chart_data, ensure_ascii=False)

    return render(request, "ingestion/chart.html", {
        "source": source,
        "chart_data": chart_data_json
    })

def user_data_summary_view(request):
    user = request.user
    sources = DataSource.objects.filter(owner=user).order_by("-created_at")

    all_data = []
    for src in sources:
        schema = src.internal_schema
        if schema:
            tables_info = get_schema_info(**src.get_db_credentials())
            all_data.append({
                "file": src.name,
                "schema": schema,
                "tables": tables_info
            })

    return render(request, "ingestion/user_summary.html", {"all_data": all_data})

@require_POST
def delete_source(request, source_id):
    # Obtenemos el dataset y su DataSource relacionado
    dataset = get_object_or_404(UploadedDataset, id=source_id)
    schema = dataset.source.internal_schema
    table = dataset.source.internal_table

    with connection.cursor() as cursor:
        # 1) Borrar primero los registros dependientes
        cursor.execute("DELETE FROM ingestion_diagrama WHERE data_source_id = %s;", [dataset.source.id])
        cursor.execute("DELETE FROM ingestion_uploadeddataset WHERE id = %s;", [dataset.id])
        cursor.execute("DELETE FROM ingestion_externalconnection WHERE source_id = %s;", [dataset.source.id])

        # 2) Ahora sí podemos borrar el DataSource
        cursor.execute("DELETE FROM ingestion_datasource WHERE id = %s;", [dataset.source.id])

        # 3) Borrar el esquema asociado (tablas dinámicas del archivo)
        cursor.execute(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE;')

    return redirect("dashboard")

def download_schema(request, source_id):
    dataset = get_object_or_404(UploadedDataset, id=source_id)
    schema = dataset.source.internal_schema  # el schema del dataset

    db = settings.DATABASES["default"]

    # Crear archivo temporal
    tmp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".sql")
    tmp_file.close()

    # Comando pg_dump
    cmd = [
        "C:/Program Files/PostgreSQL/17/bin/pg_dump.exe",
        "-h", db["HOST"],
        "-p", str(db["PORT"]),
        "-U", db["USER"],
        "-d", db["NAME"],
        "--no-owner",
        "--schema", schema,
        "-f", tmp_file.name
    ]

    env = os.environ.copy()
    env["PGPASSWORD"] = db["PASSWORD"]

    # Ejecutar pg_dump
    subprocess.run(cmd, env=env, check=True)

    # Devolver archivo como descarga
    return FileResponse(open(tmp_file.name, "rb"), as_attachment=True, filename=f"{schema}.sql")

def prueba_view(request):
    sql = None
    grafico = None
    datos = []

    # 📌 Traer nombre de archivo + schema del usuario
    user = request.user
    sources = DataSource.objects.filter(owner=user).order_by("-created_at")

    archivos = []
    for src in sources:
        if src.internal_schema:
            archivos.append({
                "file": src.name,       # nombre del archivo
                "schema": src.internal_schema  # schema real en DB
            })

    if request.method == "POST":
        schema_seleccionado = request.POST.get("schema")  # el schema real
        pregunta = request.POST.get("pregunta")

        # ✅ Obtener esquema de tablas para ese schema real
        esquema = get_schema_info(schema_seleccionado)

        # ✅ Generar SQL y tipo gráfico con Gemini
        sql, grafico = generar_consulta_y_grafico(esquema, pregunta)

        # ✅ Ejecutar SQL si existe
        if sql:
            with connection.cursor() as cursor:
                cursor.execute(f"SET search_path TO {schema_seleccionado}")
                cursor.execute(sql)
                columnas = [col[0] for col in cursor.description]
                datos = [dict(zip(columnas, fila)) for fila in cursor.fetchall()]

    return render(request, "ingestion/prueba.html", {
        "archivos": archivos,   # 📌 Enviamos la lista al template
        "sql": sql,
        "grafico": grafico,
        "datos": datos
    })

def obtener_esquema_bd(schema_name):
    """Obtiene tablas y columnas del esquema seleccionado"""
    with connection.cursor() as cursor:
        cursor.execute("""
            SELECT table_name, column_name
            FROM information_schema.columns
            WHERE table_schema = %s
            ORDER BY table_name, ordinal_position
        """, [schema_name])
        filas = cursor.fetchall()

    esquema = {}
    for tabla, col in filas:
        esquema.setdefault(tabla, []).append(col)

    return esquema

def dashboard(request):
    engine = get_engine()
    with engine.begin() as conn:
        res = conn.execute(sa_text("""
            SELECT schema_name
            FROM information_schema.schemata
            WHERE schema_name LIKE 'user_%'
            ORDER BY schema_name
        """))
        schemas = [r[0] for r in res.fetchall()]

    return render(request, "dashboard.html", {"schemas": schemas})

@csrf_exempt
def prueba_chat_view(request):
    if request.method != "POST":
        user = request.user
        sources = DataSource.objects.filter(owner=user).order_by("-created_at")
        archivos = [{"file": s.name, "schema": s.internal_schema} for s in sources if s.internal_schema]
        from django.shortcuts import render
        return render(request, "ingestion/prueba.html", {"archivos": archivos})

    body = json.loads(request.body or "{}")
    schema = body.get("schema")
    mensaje = body.get("mensaje", "").strip()

    esquema_full = get_schema_info(schema)         # trae columns/rows/preview...
    esquema_reducido = reduce_schema(esquema_full) # SOLO {tabla: [columnas]}

    sql, grafico, respuesta = generar_consulta_y_grafico(esquema_reducido, mensaje)

    columns, datos, error = [], [], None
    if sql and schema:
        try:
            with connection.cursor() as cursor:
                cursor.execute(f'SET search_path TO "{schema}"')
                cursor.execute(sql)
                columns = [c[0] for c in cursor.description] if cursor.description else []
                rows = cursor.fetchall()
                datos = [list(r) for r in rows]
        except Exception as e:
            error = str(e)

    return JsonResponse(
        {
            "respuesta": respuesta or "",
            "sql": sql or "",
            "grafico": grafico or "bar",
            "columns": columns,
            "datos": datos,
            "error": error,
        },
        encoder=DjangoJSONEncoder,
        json_dumps_params={"ensure_ascii": False}
    )

@csrf_exempt
def enviar_email_view(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Método no permitido"})
    
    try:
        print("📧 Iniciando envío de email...")
        data = json.loads(request.body)
        destinatario = data.get('destinatario')
        asunto = data.get('asunto')
        mensaje = data.get('mensaje')
        attachment_data = data.get('attachment')
        file_name = data.get('fileName')
        
        print(f"📧 Destinatario: {destinatario}")
        print(f"📧 Asunto: {asunto}")
        print(f"📧 Configuración EMAIL_HOST_USER: {settings.EMAIL_HOST_USER}")
        
        if not destinatario or not asunto:
            print("❌ Faltan datos requeridos")
            return JsonResponse({"success": False, "error": "Email y asunto requeridos"})
        
        if '@' not in destinatario:
            print("❌ Email inválido")
            return JsonResponse({"success": False, "error": "Email inválido"})
        
        # Crear email
        print("📧 Creando mensaje de email...")
        email = EmailMessage(
            subject=asunto,
            body=mensaje,
            from_email=settings.EMAIL_HOST_USER,
            to=[destinatario],
        )
        
        # Agregar adjunto PDF
        if attachment_data and file_name:
            print(f"📎 Procesando adjunto: {file_name}")
            header, encoded = attachment_data.split(',', 1)
            file_data = base64.b64decode(encoded)
            email.attach(file_name, file_data, 'application/pdf')
            print("✅ Adjunto agregado")
        
        # Enviar
        print("📤 Enviando email...")
        result = email.send()
        print(f"✅ Email enviado. Resultado: {result}")
        
        return JsonResponse({"success": True, "message": f"Email enviado a {destinatario}"})
        
    except Exception as e:
        print(f"❌ Error completo: {str(e)}")
        import traceback
        traceback.print_exc()
        return JsonResponse({"success": False, "error": str(e)})

# ========================================
# NUEVAS VISTAS PARA DIAGRAMAS
# ========================================

@login_required
def dashboard_view(request, source_id):
    """
    Vista principal del dashboard para un archivo específico.
    Muestra diagramas automáticos + chat integrado.
    """
    source = get_object_or_404(DataSource, id=source_id, owner=request.user)
    diagramas = obtener_diagramas_por_archivo(source)
    
    context = {
        'source': source,
        'diagramas': diagramas,
        'diagramas_automaticos': diagramas.filter(source_type=Diagrama.AUTO),
        'diagramas_chat': diagramas.filter(source_type=Diagrama.CHAT),
    }
    
    return render(request, 'ingestion/dashboard_view.html', context)

@csrf_exempt

@login_required
def chat_integrado_view(request):
    """
    Chat integrado específico para un archivo.
    Espera JSON:
      {
        "source_id": <int>,
        "mensaje": <str>,
        "db": {                     # (opcional) override de credenciales
          "host": "...",
          "port": 5432,
          "database": "...",
          "user": "...",
          "password": "...",
          "schema": "public"
        }
      }
    """
    if request.method != "POST":
        return JsonResponse({"error": "Método no permitido"}, status=405)

    try:
        try:
            payload = json.loads(request.body or "{}")
        except JSONDecodeError:
            return JsonResponse({"error": "Body debe ser JSON válido"}, status=400)

        source_id = payload.get("source_id")
        mensaje = (payload.get("mensaje") or "").strip()
        override_creds = payload.get("db") or {}

        if not source_id or not isinstance(source_id, int):
            return JsonResponse({"error": "source_id requerido (int)"}, status=400)

        if not mensaje:
            return JsonResponse({"error": "mensaje requerido"}, status=400)

        source = get_object_or_404(DataSource, id=source_id, owner=request.user)

        # Router IA: SQL vs ML
        diagrama, error, ask, ml_payload = generar_chat_chart_o_prediccion(
            source,
            mensaje,
            **override_creds  # permite forzar host/port/db/user/password/schema si llega del front
        )

        # Errores de negocio (conexión, SQL, ML, etc.)
        if error:
            return JsonResponse({"success": False, "error": error}, status=200)

        # Falta de contexto: pregunta de seguimiento
        if ask:
            return JsonResponse({"success": True, "ask": ask}, status=200)

        # Modo ML
        if ml_payload:
            return JsonResponse({
                "success": True,
                "mode": "ml_predict",
                "ml": ml_payload
            }, status=200)

        # Modo SQL/Chart (diagrama temporal; NO guardado aún)
        if diagrama:
            return JsonResponse({
                "success": True,
                "mode": "sql_chart",
                "diagrama": {
                    "title": diagrama.title,
                    "description": diagrama.description,
                    "chart_type": diagrama.chart_type,
                    "chart_data": diagrama.chart_data,
                    "sql_query": diagrama.sql_query
                }
            }, status=200)

        # Nada que mostrar
        return JsonResponse({"success": False, "error": "No recibí contenido para mostrar."}, status=200)

    except Exception as e:
        # Log detallado opcional
        # import traceback; traceback.print_exc()
        return JsonResponse({"error": str(e)}, status=500)

@csrf_exempt
@login_required
def guardar_diagrama_view(request):
    """
    Guarda un diagrama generado por el chat.
    """
    if request.method != "POST":
        return JsonResponse({"error": "Método no permitido"}, status=405)
    
    try:
        data = json.loads(request.body)
        source_id = data.get('source_id')
        title = data.get('title')
        description = data.get('description')
        chart_type = data.get('chart_type')
        chart_data = data.get('chart_data')
        sql_query = data.get('sql_query')
        
        if not all([source_id, title, chart_type, chart_data, sql_query]):
            return JsonResponse({"error": "Datos incompletos"}, status=400)
        
        source = get_object_or_404(DataSource, id=source_id, owner=request.user)
        
        # Crear y guardar diagrama
        diagrama = Diagrama.objects.create(
            data_source=source,
            owner=request.user,
            title=title,
            description=description,
            chart_type=chart_type,
            source_type=Diagrama.CHAT,
            sql_query=sql_query,
            chart_data=chart_data,
            order=source.diagramas_count
        )
        
        return JsonResponse({
            "success": True,
            "diagrama_id": diagrama.id,
            "message": "Diagrama guardado exitosamente"
        })
        
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)

@login_required
def listar_diagramas_view(request, source_id):
    """
    Lista todos los diagramas de un archivo específico.
    """
    source = get_object_or_404(DataSource, id=source_id, owner=request.user)
    diagramas = obtener_diagramas_por_archivo(source)
    
    return render(request, 'ingestion/listar_diagramas.html', {
        'source': source,
        'diagramas': diagramas
    })

@require_POST
@login_required
def eliminar_diagrama_view(request, diagrama_id):
    """
    Elimina un diagrama específico.
    """
    diagrama = get_object_or_404(Diagrama, id=diagrama_id, owner=request.user)
    source_id = diagrama.data_source.id
    diagrama.delete()
    
    return redirect('ingestion:dashboard_view', source_id=source_id)

@login_required
def actualizar_diagrama_view(request, diagrama_id):
    """
    Re-ejecuta la consulta de un diagrama para actualizar datos.
    """
    diagrama = get_object_or_404(Diagrama, id=diagrama_id, owner=request.user)
    
    try:
        chart_data = diagrama.execute_query()
        return JsonResponse({
            "success": True,
            "chart_data": chart_data,
            "message": "Diagrama actualizado"
        })
    except Exception as e:
        return JsonResponse({
            "success": False,
            "error": str(e)
        }, status=500)

@login_required
def analyze_chart_view(request):
    context = {"result": None, "error": None}
    if request.method == "POST" and request.FILES.get("image"):
        image_file = request.FILES["image"]
        if image_file.content_type not in ("image/png", "image/jpeg", "image/jpg", "image/webp"):
            context["error"] = "Formato no soportado. Sube PNG/JPG/WEBP."
        elif image_file.size > 5 * 1024 * 1024:
            context["error"] = "La imagen supera 5MB."
        else:
            tmp_path = default_storage.save(f"uploads/{uuid.uuid4().hex}_{image_file.name}", image_file)
            abs_path = default_storage.path(tmp_path)
            try:
                result = analyze_chart_image(abs_path)
                context["result"] = result
            except Exception as e:
                context["error"] = str(e)
    return render(request, "ingestion/analyze_chart.html", context)


@login_required
def drag_drop_view(request, source_id):
    """
    Página principal Drag & Drop para un DataSource.
    """
    source = get_object_or_404(DataSource, id=source_id, owner=request.user)
    # Puedes pasar info básica, el JS pedirá el esquema vía API
    return render(request, "ingestion/drag_drop.html", {
        "source": source
    })

@login_required
def dragdrop_schema_api(request, source_id):
    """
    Devuelve {tabla: {columns:[...], rows:int}} del schema del source.
    """
    source = get_object_or_404(DataSource, id=source_id, owner=request.user)
    if not source.internal_schema:
        return JsonResponse({"error": "Este archivo no tiene schema interno"}, status=400)

    esquema = get_schema_info(source.internal_schema)
    # Reducir un poco el payload
    out = {
        "schema": source.internal_schema,
        "tables": {
            t: {"columns": info.get("columns", []), "rows": info.get("rows", 0)}
            for t, info in esquema.items()
        }
    }
    return JsonResponse(out)

@csrf_exempt
@login_required
def dragdrop_run_api(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "Método no permitido"}, status=405)

    try:
        payload = json.loads(request.body or "{}")
        source_id = payload.get("source_id")
        chart_type = (payload.get("chart_type") or "").lower()
        table = payload.get("table") or ""
        legend = payload.get("legend") or ""
        value = payload.get("value") or None
        agg = (payload.get("agg") or "COUNT").upper()
        limit = int(payload.get("limit") or 10)

        src = get_object_or_404(DataSource, id=source_id, owner=request.user)
        schema = src.internal_schema
        if not schema or not table or not legend:
            return JsonResponse({"success": False, "error": "Faltan parámetros (schema/tabla/leyenda)."}, status=400)

        # Saneos básicos (solo para nombres; igualmente usamos comillas dobles)
        table = sanitize_identifier(table)

        # Armar expresión de agregación
        if agg == "COUNT":
            agg_expr = "COUNT(*)"
            value_label = "count"
            value_not_null = ""
        else:
            if not value:
                return JsonResponse({"success": False, "error": "Para SUM/AVG/MIN/MAX se requiere 'value'."}, status=400)
            agg_expr = f'{agg}("{value}")'
            value_label = f"{agg.lower()}({value})"
            value_not_null = f' AND "{value}" IS NOT NULL'

        engine = get_engine(**src.get_db_credentials())
        with engine.begin() as conn:

            if chart_type == "scatter":
                # Detectar tipo de la columna leyenda (X)
                dtype = conn.execute(sa_text("""
                    SELECT data_type 
                    FROM information_schema.columns
                    WHERE table_schema = :s AND table_name = :t AND column_name = :c
                """), {"s": schema, "t": table, "c": legend}).scalar() or ""

                numeric_types = {"integer","bigint","numeric","real","double precision","decimal","smallint"}
                if dtype.lower() in numeric_types:
                    x_expr = f'"{legend}"'
                else:
                    # Casteo seguro si la leyenda es texto con números
                    x_expr = f"""
                        CASE 
                        WHEN "{legend}" ~ '^[+-]?[0-9]*\\.?[0-9]+$' THEN ("{legend}")::numeric 
                        ELSE NULL 
                        END
                    """

                # Para SUM/AVG/MIN/MAX requerimos value; para COUNT no
                if agg != "COUNT" and not value:
                    return JsonResponse({"success": False, "error": "Para SUM/AVG/MIN/MAX se requiere 'value'."}, status=400)

                # Construir SQL con CTE y filtrar alias en la capa externa
                sql = f"""
                    WITH pts AS (
                        SELECT
                            {x_expr} AS x,
                            {agg_expr} AS y
                        FROM "{schema}"."{table}"
                        WHERE "{legend}" IS NOT NULL {value_not_null}
                        GROUP BY 1
                    )
                    SELECT x, y
                    FROM pts
                    WHERE x IS NOT NULL
                    ORDER BY x
                    LIMIT :lim
                """

                rows = conn.execute(sa_text(sql), {"lim": limit}).fetchall()
                points = []
                for r in rows:
                    x, y = r[0], r[1]
                    if x is None or y is None:
                        continue
                    try:
                        points.append({"x": float(x), "y": float(y)})
                    except Exception:
                        continue

                if not points:
                    return JsonResponse({"success": False, "error": "Sin puntos válidos para scatter (¿leyenda numérica?)"})

                chart_data = {
                    "datasets": [{
                        "label": value_label,  # p.ej. 'count' o 'sum(col)'
                        "data": points
                    }]
                }
                return JsonResponse({"success": True, "chart_type": "scatter", "chart_data": chart_data})
                    

            # ---- Resto de tipos (bar/line/pie/doughnut/radar): labels + values ----
            sql = f"""
                SELECT "{legend}" AS label, {agg_expr} AS value
                FROM "{schema}"."{table}"
                WHERE "{legend}" IS NOT NULL {value_not_null}
                GROUP BY "{legend}"
                ORDER BY value DESC
                LIMIT :lim
            """
            rows = conn.execute(sa_text(sql), {"lim": limit}).fetchall()
            labels = [str(r[0]) for r in rows]
            values = [float(r[1]) if r[1] is not None else 0 for r in rows]

            chart_data = {
                "labels": labels,
                "datasets": [{
                    "label": value_label,
                    "data": values
                }]
            }
            # radar usa este mismo formato
            return JsonResponse({"success": True, "chart_type": chart_type, "chart_data": chart_data, "sql": sql,})

    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)
    


def descargar_excel_view(request, schema):
    """Genera un archivo Excel con formato correcto para datetime y booleanos."""
    info = get_schema_info(schema)

    output = io.BytesIO()
    writer = pd.ExcelWriter(output, engine="openpyxl", datetime_format="YYYY-MM-DD HH:MM:SS")

    for table, data in info.items():
        df = pd.DataFrame(data["preview"], columns=data["columns"])

        if df.empty:
            df = pd.DataFrame(columns=data["columns"])

        # Convertir timestamps con zona horaria → naive
        for col in df.select_dtypes(include=["datetimetz"]).columns:
            df[col] = df[col].dt.tz_localize(None)

        # Convertir booleanos a "Sí"/"No"
        for col in df.select_dtypes(include=["bool"]).columns:
            df[col] = df[col].map({True: "Sí", False: "No"})

        # Convertir objetos raros a texto
        for col in df.select_dtypes(include=["object"]).columns:
            df[col] = df[col].astype(str)

        # Guardar en la hoja de Excel
        df.to_excel(writer, sheet_name=table[:31], index=False)

        # Ajustar ancho de columnas
        ws = writer.sheets[table[:31]]
        for i, col in enumerate(df.columns, start=1):
            max_len = max(df[col].astype(str).map(len).max(), len(col)) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max_len, 50)

    writer.close()
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{schema}.xlsx"'
    return response   